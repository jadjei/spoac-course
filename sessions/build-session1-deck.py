#!/usr/bin/env python3
"""Build the Session 1 deck in the same visual system as spo-academy-onboarding.pptx.

Palette, fonts and geometry were read back off the onboarding deck rather than
guessed: dark 0F2027, panel F7FAFB, accent 0891B2, ink 1F2937, muted 6B7280,
Poppins for headings, Noto Sans for body, 13.33x7.5in slides, logo top right.

Usage:
    python3 build-session1-deck.py <onboarding.pptx> <out.pptx> <logo.png>
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
import copy, re, sys

if len(sys.argv) != 4:
    sys.exit("usage: build-session1-deck.py <onboarding.pptx> <out.pptx> <logo.png>")
SRC, OUT, LOGO = sys.argv[1:4]

DARK, PANEL, CYAN, CYAN_BRIGHT = "0F2027", "F7FAFB", "0891B2", "06B6D4"
INK, MUTED, ON_DARK, WHITE = "1F2937", "6B7280", "B9C6CC", "FFFFFF"
HEAD, BODY = "Poppins", "Noto Sans"

def rgb(h): return RGBColor.from_string(h)

prs = Presentation(SRC)
for sid in list(prs.slides._sldIdLst):      # start from the template, empty of slides
    prs.part.drop_rel(sid.rId); prs.slides._sldIdLst.remove(sid)
BLANK = prs.slide_masters[0].slide_layouts[0]

def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    for sh in list(s.shapes):               # layout placeholders are not wanted
        sh._element.getparent().remove(sh._element)
    # Full-bleed rectangle rather than <p:bg>: Google Slides silently drops a
    # slide-level background on pptx import, which turns a dark slide white and
    # hides every piece of white text on it. A shape survives the round trip.
    if bg != WHITE:
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        r.fill.solid(); r.fill.fore_color.rgb = rgb(bg)
        r.line.fill.background(); r.shadow.inherit = False
        if r.has_text_frame: r.text_frame.clear()
    return s

def text(s, x, y, w, h, body, size, colour, font=BODY, bold=False, space=None, align=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(body if isinstance(body, list) else [body]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if space: p.space_after = Pt(space)
        if align: p.alignment = align
        r = p.add_run(); r.text = line
        r.font.name = font; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = rgb(colour)
    return tb

def panel(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill); sp.line.fill.background()
    sp.shadow.inherit = False
    if sp.has_text_frame: sp.text_frame.clear()
    return sp

def head(s, eyebrow, title):
    text(s, 0.7, 0.45, 8.0, 0.3, eyebrow.upper(), 11, CYAN, BODY, True)
    text(s, 0.7, 0.78, 10.8, 0.9, title, 34, INK, HEAD, True)
    s.shapes.add_picture(LOGO, Inches(12.05), Inches(0.42), Inches(0.58), Inches(0.58))

def chip_row(s, y, num, label, detail, x=0.7, wide=11.9):
    panel(s, x, y, 0.62, 0.62, CYAN, MSO_SHAPE.OVAL)
    text(s, x, y + 0.13, 0.62, 0.4, num, 13, WHITE, HEAD, True, align=PP_ALIGN.CENTER)
    text(s, x + 0.85, y + 0.02, 4.2, 0.4, label, 17, INK, HEAD, True)
    if detail: text(s, x + 0.85, y + 0.36, wide - 1.2, 0.5, detail, 13, MUTED)

# 1 - title
s = slide(DARK)
s.shapes.add_picture(LOGO, Inches(0.7), Inches(1.15), Inches(0.95), Inches(0.95))
text(s, 0.7, 2.25, 11.0, 1.0, "SPO Academy", 54, WHITE, HEAD, True)
text(s, 0.7, 3.3, 11.0, 0.6, "Session 1 - Orientation and Our Changing World", 28, CYAN_BRIGHT, HEAD)
text(s, 0.7, 4.25, 11.0, 0.4, "Monday 28 September 2026  ·  16:00 UTC  ·  Jonathan and Harry", 14, ON_DARK)

# 2 - where the cohort is
s = slide(); head(s, "Where we are", "Some of you are ahead. Some are not on yet.")
panel(s, 0.7, 2.2, 11.9, 1.2, PANEL)
text(s, 1.15, 2.4, 2.4, 0.4, "Already working", 19, CYAN, HEAD, True)
text(s, 3.7, 2.44, 8.5, 0.8, "Assignments submitted and accepted on Andamio. Modules 000 and 100 under way.", 14, INK)
panel(s, 0.7, 3.75, 11.9, 1.2, PANEL)
text(s, 1.15, 3.95, 2.4, 0.4, "Not on yet", 19, CYAN, HEAD, True)
text(s, 3.7, 3.99, 8.5, 0.8, "No access token minted. That is a five minute job and we are doing it on this call.", 14, INK)
text(s, 0.7, 6.1, 11.9, 0.4, "Nobody leaves this session blocked. That is the one promise for tonight.", 13, MUTED)

# 3 - token clinic
s = slide(); head(s, "First thing", "Access token clinic, live, right now")
chip_row(s, 2.2, "1", "A mainnet wallet", "A fresh, separate wallet is fine. It holds your credential, so it must be one you control.")
chip_row(s, 3.3, "2", "Mint the access token", "Small amount of ADA on mainnet. Follow along on the shared screen.")
chip_row(s, 4.4, "3", "Fluid7 refunds the cost", "Send me the transaction and the ADA comes back. Nobody is out of pocket for being here.")
text(s, 0.7, 6.1, 11.9, 0.4, "Stuck after tonight? Post in the cohort Discord rather than going quiet.", 13, CYAN, BODY, True)

# 4 - module 000
s = slide(); head(s, "Module 000", "What this course is, and is not")
panel(s, 0.7, 2.2, 11.9, 1.2, PANEL)
text(s, 1.15, 2.4, 2.2, 0.4, "It is", 19, CYAN, HEAD, True)
text(s, 3.4, 2.44, 8.8, 0.8, "The business and marketing side of running a pool. Who it is for, why it exists, how it grows.", 14, INK)
panel(s, 0.7, 3.75, 11.9, 1.2, PANEL)
text(s, 1.15, 3.95, 2.2, 0.4, "It is not", 19, CYAN, HEAD, True)
text(s, 3.4, 3.99, 8.8, 0.8, "A technical course. Technical questions still get answered, they are simply not the subject.", 14, INK)
text(s, 0.7, 6.1, 11.9, 0.4, "Eight modules  ·  two to three hours a week  ·  an on-chain credential you keep", 13, MUTED)

# 5 - the premise
s = slide(); head(s, "The premise", "Technically strong, commercially invisible")
text(s, 0.7, 2.15, 7.3, 1.6, "You can hit every block, keep uptime clean and rotate keys on time, and still watch delegation go elsewhere.", 16, INK)
text(s, 0.7, 3.9, 7.3, 1.2, "That gap is not a technical problem. So no amount of technical work will close it.", 16, INK)
panel(s, 8.5, 2.1, 4.1, 2.9, DARK)
text(s, 8.85, 2.65, 3.4, 1.8, ["Nobody outside", "the operator can", "say why your", "pool exists"], 22, WHITE, HEAD, True, space=2)

# 6 - module 100
s = slide(); head(s, "Module 100", "Our changing world")
chip_row(s, 2.2, "1", "Undelegated stake", "A large amount of ADA has never been delegated to anything at all.")
chip_row(s, 3.3, "2", "Different pitches", "Different projects and holders need different value propositions, not one general pitch.")
chip_row(s, 4.4, "3", "A crowded field", "The chain parameters target 500 pools. That is what the protocol aims at, not a count of pools already at optimal size.")
text(s, 0.7, 6.1, 11.9, 0.4, "This is the discussion slide. Bring your own experience, especially where it contradicts mine.", 13, MUTED)

# 7 - activation argument
s = slide(); head(s, "The argument worth having", "Activation, not competition")
text(s, 0.7, 2.15, 7.3, 1.8, "If a large share of ADA has never been staked at all, then growth may be about activating people who have never delegated, rather than taking stake from pools that already hold it.", 16, INK)
text(s, 0.7, 4.2, 7.3, 0.8, "Those are two completely different marketing jobs.", 16, CYAN, BODY, True)
panel(s, 8.5, 2.1, 4.1, 2.9, DARK)
text(s, 8.85, 2.75, 3.4, 1.6, ["Have you ever", "marketed to", "someone who has", "never staked?"], 22, WHITE, HEAD, True, space=2)

# 8 - your turn
s = slide(); head(s, "Your turn", "Who is your pool for?")
text(s, 0.7, 2.3, 11.9, 1.0, "One sentence each. Not the copy on your website.", 20, INK)
panel(s, 0.7, 3.4, 11.9, 1.5, PANEL)
text(s, 1.15, 3.75, 11.0, 0.9, "Say it out loud, we write them all down, and they become the raw material for module 300.", 16, INK)
text(s, 0.7, 6.1, 11.9, 0.4, "If it is hard to answer, that is the finding, not a failure.", 13, MUTED)

# 9 - session 2 timing
s = slide(); head(s, "Housekeeping", "Session 2 needs a new slot")
text(s, 0.7, 2.15, 11.9, 0.9, "The scheduled date clashes with Token2049, and enough of you are affected that moving it is the right call.", 16, INK)
panel(s, 0.7, 3.2, 5.8, 1.9, PANEL)
text(s, 1.1, 3.5, 5.0, 0.4, "What we need", 19, CYAN, HEAD, True)
text(s, 1.1, 3.95, 5.0, 1.0, "A date and a time that works for the room, agreed now rather than over email.", 14, INK)
panel(s, 6.8, 3.2, 5.8, 1.9, DARK)
text(s, 7.2, 3.5, 5.0, 0.4, "Sessions 3 and 4", 19, CYAN_BRIGHT, HEAD, True)
text(s, 7.2, 3.95, 5.0, 1.0, "12 and 19 October, 16:00 UTC, unchanged unless you tell me otherwise.", 14, ON_DARK)

# 10 - before next time
s = slide(); head(s, "Before next time", "Three things")
chip_row(s, 2.3, "1", "Module 200", "Your Journey Toolkit. The frameworks the rest of the course runs on.")
chip_row(s, 3.5, "2", "Its assignment", "Fifteen to thirty minutes. These accumulate into a business plan you keep.")
chip_row(s, 4.7, "3", "Start 300 if you can", "Your Why. It is the one that takes longest to sit with.")

# 11 - recording
s = slide(); head(s, "Catch-up", "Every session is recorded from tonight")
text(s, 0.7, 2.2, 11.9, 1.0, "Recordings and written summaries go up on the cohort pages, so missing a week costs you nothing.", 16, INK)
panel(s, 0.7, 3.5, 11.9, 1.3, PANEL)
text(s, 1.15, 3.85, 11.0, 0.7, "Zoom announces recording to everyone present, every time it is running.", 14, INK)
text(s, 0.7, 6.1, 11.9, 0.4, "The onboarding recording lost its first fifteen minutes to a storage limit. That is fixed.", 13, MUTED)

# 12 - close
s = slide(DARK)
text(s, 0.7, 2.6, 11.0, 0.9, "Nobody leaves blocked", 44, WHITE, HEAD, True)
text(s, 0.7, 3.7, 11.0, 1.2, "If your token is not minted by the end of this call, we book five minutes this week and finish it together.", 18, ON_DARK)
text(s, 0.7, 5.3, 11.0, 0.4, "Next session: date to be agreed tonight  ·  same Zoom link every time", 14, CYAN_BRIGHT)

prs.save(OUT)
print("wrote", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
