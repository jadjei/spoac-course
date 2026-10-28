#!/usr/bin/env python3
"""Regenerate `plan/course-plan.yaml` from the current SPOAC Course Planner.

The planner is authoritative over both the live Andamio structure and anything
previously scaffolded here. This script is the only sanctioned way to rebuild the plan
— edit the planner, re-run this, never hand-edit the yaml.

Sheet shape (`SPO Academy Modules`, header row 7):

    Module | Old | Module Title | Status | Element Type
    [General Mod Overview Filename] | SPO Overview FIlename | Snagging
    Title | Overview | Student Learning Target

Columns are resolved by their header text, not by position. The planner has been
reshaped three times in two days, once by deleting the General filename column, which
shifted everything after it one to the left. Reading by position turned the SPO decks
into General ones silently, and the count was the only thing that gave it away. A header
that goes missing raises here instead.

Row kinds:
  * `A` matches ^\\d00$          -> a module (000…700)
  * `A` set, e.g. 100.1 / 301.4b -> an element of the current module
  * `A` blank but the row has    -> a continuation of the element above; its Title /
    content                         Overview / SLTs enrich that element rather than
                                    becoming a new one (the planner uses these rows for
                                    per-deck detail and for SLT lists)

Column F/G hyperlinks are resolved to the local deck paths, so each element carries the
deck its narration comes from. Precedence is recorded per element as `source`:
`spo` (authoritative) > `general` (fallback, module 100 only) > none.

Not every element is sourced from a deck. The planner now records the delivered artefact
directly in the same columns for the two types that never had one:

    video       a YouTube link          -> `video_url`
    assignment  a Playbook Builder link -> `playbook_url`

Those are captured verbatim rather than resolved to anything local, and they leave
`source` alone — it stays a statement about deck narration. What they do mean is that
the element is covered, which is what `gap-report.py` reads them for.

Usage:
    python3 plan-from-planner.py                 # rewrite plan/course-plan.yaml
    python3 plan-from-planner.py --check         # report, write nothing
"""
import argparse, glob, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
PLANNER = os.path.join(HERE, "spoac-dl/planner/current/SPOAC Course Planner.xlsx")
# newest lsjson snapshot wins — a deck added after the last snapshot is invisible otherwise
LISTING = sorted(glob.glob(os.path.join(HERE, "spoac-dl/drive-listing-*.json")))[-1]
SHARED = sorted(glob.glob(os.path.join(HERE, "spoac-dl/general-listing-*.json")))[-1]
OUT = os.path.join(HERE, "plan/course-plan.yaml")
DEFERRED = os.path.join(HERE, "plan/deferred.yaml")

COURSE_ID = "e60f29abf7a3f348bf4f51a078dae6f53cd3bc065d74632d3f67d785"
MODULE_RE = re.compile(r"^\d00$")
LINK_RE = re.compile(r'HYPERLINK\("([^"]+)"')
ID_RE = re.compile(r"/d/([A-Za-z0-9_-]{20,})")
# the delivered artefact for the two element types that never had a deck
YOUTUBE_RE = re.compile(r"youtu\.be/|youtube\.com/(?:watch|embed|live)", re.I)
PLAYBOOK_RE = re.compile(r"spo\.academy/playbook", re.I)


def planner_date():
    """Date of the planner we actually read — rclone preserves Drive's modified time."""
    import datetime
    return datetime.date.fromtimestamp(os.path.getmtime(PLANNER)).isoformat()


HEADER_ROW = 7
# header text -> the key this script uses. Matching is loose because the planner's own
# spelling is not stable ("SPO Overview FIlename").
WANTED = {
    "module": "code", "old": "old", "module title": "mtitle", "status": "status",
    "element type": "etype", "spo overview filename": "spo",
    "general mod overview filename": "general", "title": "title", "overview": "overview",
    "student learning target": "slt",
}
REQUIRED = ("code", "etype", "spo", "title", "overview", "slt")


def columns(ws):
    """Column index per key, resolved from the header row rather than assumed."""
    found = {}
    for c in range(1, ws.max_column + 1):
        h = re.sub(r"\s+", " ", str(ws.cell(HEADER_ROW, c).value or "")).strip().lower()
        if h in WANTED and WANTED[h] not in found:
            found[WANTED[h]] = c
    missing = [k for k in REQUIRED if k not in found]
    if missing:
        sys.exit(f"planner header row {HEADER_ROW} is missing: {', '.join(missing)}. "
                 f"Found: {[str(ws.cell(HEADER_ROW, c).value) for c in range(1, ws.max_column + 1)]}")
    return found


def cell(ws, r, c):
    v = ws.cell(r, c).value
    return "" if v is None else str(v).strip()


def links(ws, r, col):
    """Drive ids hyperlinked from a cell, formula-style or as a real hyperlink."""
    c = ws.cell(r, col)
    out = []
    if c.hyperlink and c.hyperlink.target:
        out.append(c.hyperlink.target)
    out += LINK_RE.findall(str(c.value or ""))
    return [m.group(1) for m in (ID_RE.search(t) for t in out) if m]


def weblinks(ws, r, col):
    """Plain http(s) links in a cell — a real hyperlink, a HYPERLINK() formula, or the
    cell text itself, which is how the planner most often holds them."""
    c = ws.cell(r, col)
    out = []
    if c.hyperlink and c.hyperlink.target:
        out.append(c.hyperlink.target)
    out += LINK_RE.findall(str(c.value or ""))
    out.append(str(c.value or "").strip())
    return [u for u in out if u.lower().startswith("http")]


def classify(urls):
    """(video, playbook) — the most specific link of each kind, ignoring anything else.

    Specificity matters because the stored hyperlink target loses the fragment: the
    Playbook cells read `…/playbook#module-000` on the face and `…/playbook` in the
    hyperlink, and the anchor is the half that says which module it is.
    """
    def best(rx):
        hits = [u for u in urls if rx.search(u)]
        return max(hits, key=lambda u: ("#" in u, len(u)), default="")
    return best(YOUTUBE_RE), best(PLAYBOOK_RE)


def pick_title(title, mtitle, code):
    """Planner hygiene: some rows (100.2, 100.3) hold a whole paragraph in the Title
    cell. Anything that long is prose, not a title — fall back to the Module Title
    column and keep the paragraph as the overview."""
    if title and len(title) <= 80:
        return title, ""
    if mtitle:
        return mtitle, title
    return (title[:70] if title else code), title


def split_slts(text):
    if not text:
        return []
    parts = re.split(r"\n\s*-\s*|^\s*-\s*", text.strip())
    return [re.sub(r"\s+", " ", p).strip() for p in parts if p and p.strip()]


def yaml_str(s):
    s = re.sub(r"\s+", " ", str(s)).strip()
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"') + '"'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args()
    import openpyxl

    # deferral is recorded locally, never in Harry's planner — see plan/deferred.yaml
    defer = {}
    if os.path.exists(DEFERRED):
        import yaml as _y
        for d in (_y.safe_load(open(DEFERRED)) or {}).get("deferred") or []:
            defer[str(d["code"])] = d.get("reason", "")

    deck = {}
    for path, prefix in ((LISTING, "spoac-dl/spo-decks"), (SHARED, "spoac-dl/general-decks")):
        if os.path.exists(path):
            for f in json.load(open(path)):
                deck[f["ID"]] = f"{prefix}/{f['Path']}"

    ws = openpyxl.load_workbook(PLANNER)["SPO Academy Modules"]
    col = columns(ws)
    modules, mod, el = [], None, None

    for r in range(8, ws.max_row + 1):
        code, mtitle = cell(ws, r, col["code"]), cell(ws, r, col["mtitle"]) if "mtitle" in col else ""
        old = cell(ws, r, col["old"]) if "old" in col else ""
        status = cell(ws, r, col["status"]) if "status" in col else ""
        etype = cell(ws, r, col["etype"]).lower()
        title, overview = cell(ws, r, col["title"]), cell(ws, r, col["overview"])
        slt = cell(ws, r, col["slt"])
        gen = [deck[i] for i in links(ws, r, col["general"]) if i in deck] if "general" in col else []
        spo = [deck[i] for i in links(ws, r, col["spo"]) if i in deck]
        # either column may hold the artefact link — 000 and 100 keep theirs in F
        artefact_cols = [col["spo"]] + ([col["general"]] if "general" in col else [])
        video, playbook = classify([u for c in artefact_cols for u in weblinks(ws, r, c)])
        # A deck named in plain text, with no hyperlink behind it and no such file in
        # Drive, reads downstream as "nothing exists" when it means "named, not built".
        named = cell(ws, r, col["spo"])
        named = "" if (spo or named.lower().startswith("http")) else named

        if MODULE_RE.match(code):
            mod = dict(code=code, old=old, title=mtitle, status=status or "wip", elements=[])
            modules.append(mod)
            el = None
        if mod is None:
            continue

        if code:  # a coded row starts an element (the module row starts its intro element)
            use_title, spilled = pick_title(title, mtitle, code)
            el = dict(code=code, old=old, type=etype or "lesson",
                      title=use_title, overview=overview or spilled,
                      slts=split_slts(slt), spo=spo, general=gen,
                      video=video, playbook=playbook, named=named, rows=[r])
            mod["elements"].append(el)
        elif el is not None and (title or overview or slt or spo or gen or etype
                                 or video or playbook or named):
            el["rows"].append(r)  # continuation row — enrich the element above
            if not el["overview"]:
                el["overview"] = overview
            elif overview and overview not in el.get("extra", []):
                el.setdefault("extra", []).append(overview)
            if not el["title"]:
                el["title"] = title
            el["slts"] += [s for s in split_slts(slt) if s not in el["slts"]]
            el["spo"] += [d for d in spo if d not in el["spo"]]
            el["general"] += [d for d in gen if d not in el["general"]]
            el["video"] = el["video"] or video
            el["playbook"] = el["playbook"] or playbook
            el["named"] = el["named"] or named

    lines = [
        "# SPO Online Business Academy — planned structure",
        "#",
        "# GENERATED by plan-from-planner.py — do not hand-edit; edit the planner and re-run.",
        f"# Source: gdwrite:Project Management/SPOAC Course Planner.xlsx "
        f"(modified {planner_date()}).",
        "#",
        "# The planner and the deck set it links are AUTHORITATIVE over the live Andamio",
        "# structure and content — Andamio gets reshaped to match, never the reverse.",
        "#",
        "# Per element, `source` records where the lesson body comes from:",
        "#   spo     — the planner-linked _SPO deck (authoritative wherever one exists)",
        "#   general — the General flavour, fallback only (broader audience: cut and re-aim)",
        "#   none    — no current-generation deck; stays a stub, do not invent content",
        "# Narration lives in each deck's speaker notes; extracted under extracted/.",
        "",
        "course:",
        f"  title: SPO Online Business Academy",
        f"  course_id: {COURSE_ID}",
        "  network: mainnet",
        "",
        "modules:",
    ]
    n_el = n_spo = n_gen = 0
    for m in modules:
        lines += [f'  - code: "{m["code"]}"']
        if m["old"]:
            lines.append(f"    old: {m['old']}")
        lines += [f"    title: {yaml_str(m['title'])}", f"    status: {yaml_str(m['status'])}",
                  "    elements:"]
        for e in m["elements"]:
            n_el += 1
            src = "spo" if e["spo"] else ("general" if e["general"] else "none")
            n_spo += bool(e["spo"])
            n_gen += bool(e["general"] and not e["spo"])
            lines.append(f'      - code: "{e["code"]}"')
            if e["old"]:
                lines.append(f"        old: {e['old']}")
            lines += [f"        type: {e['type'] or 'lesson'}",
                      f"        title: {yaml_str(e['title'] or e['code'])}",
                      f"        source: {src}"]
            if str(e["code"]) in defer:
                lines += ["        deferred: true",
                          f"        deferred_reason: {yaml_str(defer[str(e['code'])])}"]
            if e["overview"]:
                lines.append(f"        overview: {yaml_str(e['overview'])}")
            for extra in e.get("extra", []):
                lines.append(f"        # also: {re.sub(chr(10), ' ', extra)[:150]}")
            if e["slts"]:
                lines.append("        slts:")
                lines += [f"          - {yaml_str(s)}" for s in e["slts"]]
            for key in ("spo", "general"):
                if e[key]:
                    lines.append(f"        {key}_decks:")
                    lines += [f"          - {yaml_str(d)}" for d in e[key]]
            if e["video"]:
                lines.append(f"        video_url: {yaml_str(e['video'])}")
            if e["playbook"]:
                lines.append(f"        playbook_url: {yaml_str(e['playbook'])}")
            if e["named"] and not e["spo"]:
                lines.append(f"        named_deck: {yaml_str(e['named'])}")
            lines.append(f"        planner_rows: {e['rows']}")
        lines.append("")

    n_def = sum(1 for m in modules for e in m["elements"] if str(e["code"]) in defer)
    n_vid = sum(1 for m in modules for e in m["elements"] if e["video"])
    n_pb = sum(1 for m in modules for e in m["elements"] if e["playbook"])
    print(f"{len(modules)} modules, {n_el} elements — {n_spo} SPO-sourced, "
          f"{n_gen} General-only, {n_el - n_spo - n_gen} with no deck"
          + (f"; {n_def} deferred" if n_def else ""))
    print(f"artefact links: {n_vid} video, {n_pb} assignment")
    missing = [c for c in defer if not any(str(e["code"]) == c
               for m in modules for e in m["elements"])]
    if missing:
        print(f"WARNING: deferred codes not found in the planner: {', '.join(sorted(missing))}")
    if args.check:
        return 0
    open(OUT, "w").write("\n".join(lines) + "\n")
    print(f"wrote {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
